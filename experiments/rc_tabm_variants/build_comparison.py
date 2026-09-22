import sys as _sys, pathlib as _pl
for _c in _pl.Path(__file__).resolve().parents:
    if (_c / 'pamp_paths.py').exists():
        _sys.path.insert(0, str(_c)); break
from pamp_paths import DATA_ROOT, REVISION_ROOT
from pathlib import Path
import json,hashlib,shutil
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl.styles import Font,PatternFill
ROOT=DATA_ROOT; PAPER=REVISION_ROOT;OUT=Path(__file__).resolve().parent
sources={}; records=[]
MET=['R2','PCC','Spearman','MSE','RMSE','MAE']
def track(p):
 p=Path(p);sources[str(p)]=hashlib.sha256(p.read_bytes()).hexdigest();return p
def csv(p):return pd.read_csv(track(p))
def js(p):return json.loads(track(p).read_text())
def add(model,metrics,source,group='matched_2766',split='test',seed='',precision='as saved',notes='',primary=False,N=None):
 d=dict(metrics);aliases={'r2':'R2','rmse':'RMSE','mae':'MAE','pearson':'PCC','spearman':'Spearman','RMSE_log2':'RMSE','MAE_log2':'MAE'}
 for k,v in aliases.items():
  if k in d:d[v]=d[k]
 rec=dict(group=group,split=split,model=model,seed=str(seed),precision=precision,N=d.get('N',N),notes=notes,source=str(source),primary=primary)
 for k in MET:rec[k]=d.get(k,np.nan)
 if pd.notna(rec['RMSE']):rec['MSE']=rec['RMSE']**2
 records.append(rec)
 return rec
# Verify actual saved split indices, not merely equal sample counts.
refpath=track(ROOT/'experiment_mean/split_indices.npz');ref=np.load(refpath);audit=[]
for d in ['experiment_mean','experiment_cls_mlp_random','experiment_compact_improvements']:
 p=track(ROOT/d/'split_indices.npz');x=np.load(p);same=all(np.array_equal(ref[k],x[k]) for k in ref.files);assert same
 audit.append(dict(experiment=d,split_match=same,evidence=str(p)))
for d in ['catapro_tabm_small','catapro_residual_condpool']:
 p=ROOT/'experiments'/d/'reports/run_manifest.json';m=js(p)
 # Preserve source evidence, metrics comparison additionally provided by saved model reports.
 audit.append(dict(experiment=d,split_match='documented reference split',evidence=str(p)))
# Final head evaluation in common FP32 reporting where available.
p=PAPER/'pamp_protein_substrate_v1/individual_head_test_metrics.csv'
for _,r in csv(p).iterrows():add(r.model,r,p,seed=42,precision=r.inference,primary=True)
p=ROOT/'experiment_cls_mlp_random/comparison_metrics.csv'
for _,r in csv(p).iterrows():
 add({'mean_cls_transformer':'CLS head','mean_interaction_mlp':'Interaction MLP','mean_cls_mlp_hybrid':'CLS+MLP Hybrid'}[r.architecture],r,p,seed=r.seed,notes='Original training-run evaluation',primary=r.architecture=='mean_cls_mlp_hybrid')
p=ROOT/'model_improvement_analysis/ensemble_diagnostics.json';d=js(p)
for name,m in d['random_validation_CPU_FP32'].items():
 add({'mean_cls_transformer':'CLS head','mean_interaction_mlp':'Interaction MLP','equal_ensemble':'Two-head average'}[name],m,p,split='val',precision='CPU FP32')
add('Two-head average',d['random_test'],p,notes='Historical numerical evaluation; canonical FP32 version listed separately')
add('Two-head average',d['official_fold_oof'],p,group='official_10fold',split='OOF')
# Compact/lowrank runs; single seed and prediction ensembles, not best-test seed selection.
labels={'esm_plain':'Compact MLP','esm_maccs':'Compact MLP + MACCS','lowrank_plain':'Low-rank interaction','lowrank_maccs':'Low-rank interaction + MACCS'}
for filename in ['esm_metrics.csv','lowrank_metrics.csv','esm_fp32_and_seed_average.csv','lowrank_fp32_and_seed_average.csv']:
 p=ROOT/'experiment_compact_improvements'/filename
 for _,r in csv(p).iterrows():
  ensemble=str(r.seed)=='equal_average_3_seeds';name=labels[r.variant]+(' (3-seed prediction ensemble)' if ensemble else '')
  add(name,r,p,seed=r.seed,precision=r.get('precision','training-run precision'),primary=ensemble)
  add(name,{'RMSE':r.val_rmse},p,split='val',seed=r.seed,precision=r.get('precision','training-run precision'),N=2766,notes='Only validation RMSE in this file; other validation metrics unavailable')
p=ROOT/'experiment_compact_improvements/ensemble_test_metrics.csv'
for _,r in csv(p).iterrows():
 if r.model=='old_CLS_MLP_plus_esm_maccs':
  add('Three-family ensemble + MACCS',r,p,seed='42/2024/3407 compact heads',primary=True)
  add('Three-family ensemble + MACCS',{'RMSE':r.val_rmse},p,split='val',N=2766)
p=PAPER/'pamp_protein_substrate_v1/prediction_metrics.csv'
for _,r in csv(p).iterrows():
 if r.model!='Two-head average':continue
 add('Two-head average',r,p,group='matched_2766' if r.dataset=='CataPro' else 'external_BRENDA',split=r.cohort if r.dataset=='CataPro' else r.cohort+'/'+r.aggregation,precision='canonical saved evaluation',primary=r.dataset=='CataPro' and r.cohort=='test')
p=ROOT/'experiments/catapro_residual_condpool/reports/metrics.json';res=js(p)
for split,m in res['metrics'].items():add('Residual Predictor',m,p,split=split,seed=42,primary=split=='test')
for split,m in res['baseline_metrics'].items():add('Original Predictor (no MACCS)',m,p,split=split,seed='two old heads + three compact seeds',primary=split=='test')
for key,name in [('tabm_mean','Mean-TabM K16'),('tabm_mean_k32','Mean-TabM K32'),('tabm_condpool','Standalone CondPool-TabM K16')]:
 p=ROOT/f'experiments/catapro_tabm_small/reports/{key}_metrics.json';d=js(p)
 for split,m in d['metrics'].items():add(name,m,p,split=split,seed=d['seed'],primary=split=='test')
for n in [2304,2471]:
 d=PAPER/f'extra_trees_quick_v1/features_{n}';cfg=js(d/'run_config.json');assert cfg['source_hashes']['split']==sources[str(refpath)]
 for split,file in [('test','test_metrics.csv'),('val','validation_metrics.json')]:
  p=d/file;m=csv(p).iloc[0] if split=='test' else js(p);add(f'Extra Trees ({n} features)',m,p,split=split,seed=42,primary=split=='test')
# External comparisons: retain model version and aggregation.
p=ROOT/'experiment_brenda_external/residual_condpool_v1/metrics.csv'
for _,r in csv(p).iterrows():add({'Residual CondPool':'Residual Predictor','Original predictor':'Original Predictor (no MACCS)','Extra Trees':'Extra Trees (2304 features)'}[r.model],r,p,group='external_BRENDA',split=r.cohort+'/'+r.unit)
p=ROOT/'experiment_brenda_external/metrics.csv'
for _,r in csv(p).iterrows():
 add({'old_cls_mlp_average':'Two-head average (historical evaluation)','improved_ensemble':'Three-family ensemble + MACCS'}[r.model],r,p,group='external_BRENDA',split=r.cohort+'/'+r.aggregation,notes='Historical external evaluation; preserve original aggregation/precision')
# All early full evaluated heads. Their 1638/1658 test sets differ from the final 2766.
for p in sorted(ROOT.glob('*/final_metrics.json')):
 d=js(p);sp=p.parent/'split_indices.npz';splits=np.load(track(sp)) if sp.exists() else None
 for split in ['train','val','test']:
  if split in d:
   n=len(splits[split+'_idx']) if splits is not None and split+'_idx' in splits else None
   add(p.parent.name,d[split],p,group='historical_different_split',split=split,N=n,notes='Not directly comparable with Residual 2766 test; see original split file')
p=ROOT/'predictor_unikp1024_ensemble_2/ensemble_metrics.json';d=js(p)
for key,split in [('validation','val'),('test','test')]:add('Early global-local two-model ensemble',d[key],p,group='historical_different_split',split=split,N=1658)
p=ROOT/'experiment_catapro_unbiased/comparison_metrics.csv'
for _,r in csv(p).iterrows():add(r.architecture,r,p,group='official_10fold',split='OOF',seed=r.seed,notes='Official fold 8 train/1 val/1 test; cannot compare directly with random-split Residual')
for folder in ['experiment_kcat_pairinfo_0618_official_dev','experiment_kcat_pairinfo_06186_official_dev']:
 p=ROOT/folder/'metrics.json';d=js(p);add(d['model'],d['dev'],p,group='official_development',split='dev',notes='independent_test=false')
p=ROOT/'experiment_reliable_kcat_v1/internal_results.json';d=js(p)
for model,groups in d.items():
 if not isinstance(groups,dict):continue
 for split,m in groups.items():
  if isinstance(m,dict) and 'point' in m:add(model,m['point'],p,group='reliability_different_protocol',split=split,notes='Different fit/selection/calibration/audit split; PCC/Spearman not saved')
p=ROOT/'experiment_reliable_kcat_v1/summary_metrics.csv'
for _,r in csv(p).iterrows():add(r.variant,r,p,group='reliability_external',split=r.cohort,notes='Deterministic substrate encoding/training protocol differs; full reliability metrics retained in source archive')
# Archive additional searches and all official folds rather than treating validation search as test evidence.
for folder,names in [('experiment_compact_improvements',['validation_search.csv','lowrank_validation_search.csv','ensemble_validation_selection.csv','summary.csv']),('experiment_catapro_unbiased',['fold_metrics.csv'])]:
 for name in names:track(ROOT/folder/name)
for folder in ['experiment_kcat_pairinfo_0618','experiment_kcat_pairinfo_0618_official_val']:
 track(ROOT/folder/'STATUS.json')
excluded=pd.DataFrame([
 dict(experiment='experiment_scmr_kcat_v0',reason='Only method review/degeneracy analysis; no completed model evaluation'),
 dict(experiment='experiments/predictor_ablation',reason='Only smoke tests found; not a full predictive benchmark'),
 dict(experiment='experiment_kcat_pairinfo_0618',reason='No final metrics.json at root; status archived; no invented final score'),
 dict(experiment='experiment_kcat_pairinfo_0618_official_val',reason='No final metrics.json at root; status archived; no invented final score'),
 dict(experiment='PAMP/A0/A1/A2/A3/B2',reason='Attack methods, not separate prediction models; excluded from predictive model ranking')])
f=pd.DataFrame(records);f['N']=pd.to_numeric(f.N,errors='coerce').astype('Int64')
assert f.R2.notna().sum()>50 and f.source.isin(sources).all()
main=f[f.primary].copy();assert main.model.is_unique and main.N.eq(2766).all()
main=main.sort_values('R2',ascending=False)
r=main[main.model.eq('Residual Predictor')].iloc[0];delta=[]
for _,v in main[~main.model.eq('Residual Predictor')].iterrows():
 row=dict(reference=v.model)
 for k in MET:row['Residual_minus_'+k]=r[k]-v[k]
 row['RMSE_reduction_percent']=100*(v.RMSE-r.RMSE)/v.RMSE;delta.append(row)
delta=pd.DataFrame(delta)
frames={'Main_test_2766':main,'Residual_differences':delta,'All_matched_records':f[f.group.eq('matched_2766')],'BRENDA_external':f[f.group.eq('external_BRENDA')],'Early_other_splits':f[f.group.eq('historical_different_split')],'Official_10fold':f[f.group.eq('official_10fold')],'Official_dev':f[f.group.eq('official_development')],'Reliability':f[f.group.str.startswith('reliability')],'Not_final_models':excluded,'Split_audit':pd.DataFrame(audit),'All_metrics':f}
for name,frame in frames.items():frame.to_csv(OUT/(name+'.csv'),index=False,float_format='%.10g')
with pd.ExcelWriter(OUT/'Residual_all_model_comparisons.xlsx',engine='openpyxl') as w:
 for name,frame in frames.items():
  frame.to_excel(w,sheet_name=name,index=False);ws=w.sheets[name];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
  for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='397B69')
  for col in ws.columns:
   letter=col[0].column_letter;header=str(col[0].value);ws.column_dimensions[letter].width=45 if header in ['model','reference','notes','source'] else 20
  for row in ws.iter_rows(min_row=2):
   for cell in row:
    if isinstance(cell.value,float):cell.number_format='0.00000'
archive=OUT/'source_tables';archive.mkdir(exist_ok=True)
for i,(path,digest) in enumerate(sources.items()):
 p=Path(path)
 if p.suffix in ['.csv','.json']:shutil.copy2(p,archive/f'{i:03d}_{p.parent.name}_{p.name}')
(OUT/'provenance.json').write_text(json.dumps(dict(source_hashes=sources,metric_count=len(f),primary_models=len(main),scope='Saved predictive results for the observed kcat model iteration lineage; no retraining or attack metrics',MSE='RMSE squared; not independently reevaluated'),indent=2))
def table(frame,cols):
 text='| '+' | '.join(cols)+' |\n| '+' | '.join(['---']*len(cols))+' |\n'
 for _,v in frame.iterrows():
  vals=[]
  for c in cols:
   x=v[c];vals.append('—' if pd.isna(x) else f'{x:.5f}' if isinstance(x,(float,np.floating)) else str(x))
  text+='| '+' | '.join(vals)+' |\n'
 return text
report='# Residual 与模型迭代历史的预测性能对比\n\n'
report+=f'本次整理 {len(f)} 条评价记录、主表 {len(main)} 个模型版本；逐种子结果、训练/验证/测试指标、外部和不同划分结果均见 Excel。来源与 SHA256 见 provenance.json，原始结果表见 source_tables。不是按时间连续改进的单一路径，包含探索失败与旁支。\n\n'
report+='## 同一固定测试集：2766 条\n\n'+table(main,['model']+MET)
report+='\n主表 Compact/Low-rank 使用三种子预测平均，不是选择测试最优种子，也不是三种子指标均值。所有单种子和历史数值精度版本均保留于 All_matched_records。缺失指标留空，MSE=RMSE²；误差单位为 log2(kcat)。\n\n'
report+='## Residual 相对各模型的差值\n\n'+table(delta,['reference','Residual_minus_R2','Residual_minus_PCC','Residual_minus_Spearman','Residual_minus_RMSE','Residual_minus_MAE','RMSE_reduction_percent'])
report+='\n## 训练与验证\n\n'+table(f[(f.group=='matched_2766')&f.model.isin(['Original Predictor (no MACCS)','Mean-TabM K16','Mean-TabM K32','Standalone CondPool-TabM K16','Residual Predictor'])&f.split.isin(['train','val'])],['model','split','N']+MET)
report+='\n## 外部评估\n\n'+table(f[(f.group=='external_BRENDA')&f.split.isin(['pair_novel/rows','sequence_novel/rows','pair_novel_all/BRENDA_rows','sequence_novel_all/BRENDA_rows'])],['model','split','N']+MET)
report+='\n外部配对未见与序列未见不是独立样本集，后者是前者子集。MACCS 集成和 Original Predictor（无 MACCS）不是同一权重版本；历史 Two-head 数值精度版本也保留命名区分，不能当成架构变化收益。\n\n'
report+='## 其他划分：不纳入主表排名\n\n'+table(f[(f.group=='historical_different_split')&(f.split=='test')],['model','N']+MET)
report+='\n'+table(f[f.group.isin(['official_10fold','official_development'])],['model','group','N']+MET)
report+='\n## 结论与解释范围\n\n在上述同划分主表中，Residual 的五项原始指标均为最优；MSE 与 RMSE 排名等价。它相对独立 CondPool 的收益较大，相对 Extra Trees 的收益较小。外部结果不支持全面超过 Extra Trees；MACCS 集成在部分外部误差指标上也更好。训练集拟合优度不代表泛化。\n\n沿用既有 UniKP 底物缓存，其位置编码沿样本轴应用的问题已在原报告记录；本表是既有流程下的结果，不是修正缓存后重训的验证。项目已多次查看测试集并迭代模型，不可描述成完全未接触的前瞻性盲测。大多数模型仅一个训练种子，主表没有进行配对显著性检验。早期 1638/1658 测试、官方十折、官方 dev、可靠性划分分别列出；不能用其 R² 与 2766 测试直接比较。\n\n'
report+='未完成/仅 smoke 的分支见 Not_final_models；没有将训练中或未保存最终指标的分支填入估计结果。Excel 保留 source 列，可逐条追溯。\n'
(OUT/'REPORT.md').write_text(report)
# Export publication-friendly static overview, no uncertainty invented.
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'pdf.fonttype':42})
plot=main.sort_values('R2');fig,axs=plt.subplots(1,2,figsize=(13,7),sharey=True)
colors=['#397B69' if x=='Residual Predictor' else '#9DAEB8' for x in plot.model]
for ax,col,title in zip(axs,['R2','RMSE'],['Test R² (higher is better)','Test RMSE (lower is better)']):
 ax.barh(plot.model,plot[col],color=colors);ax.set_title(title);ax.grid(axis='x',alpha=.2);ax.set_axisbelow(True)
 ax.spines[['top','right']].set_visible(False)
 for y,v in enumerate(plot[col]):ax.text(v+.004,y,f'{v:.4f}',va='center',fontsize=8)
 ax.set_xlim(0,plot[col].max()*1.15)
fig.suptitle('Predictive model comparisons | Same fixed test split, N = 2,766',fontsize=13)
fig.tight_layout()
for ext in ['png','pdf']:fig.savefig(OUT/f'model_comparison.{ext}',dpi=180,bbox_inches='tight')
plt.close(fig)
assert main.iloc[0].model=='Residual Predictor'
assert all((main[k].max()==r[k]) for k in ['R2','PCC','Spearman'])
assert all((main[k].min()==r[k]) for k in ['RMSE','MAE','MSE'])
writeback=pd.read_excel(OUT/'Residual_all_model_comparisons.xlsx',sheet_name='All_metrics');assert len(writeback)==len(f)
(OUT/'verification.json').write_text(json.dumps(dict(status='PASS',evaluation_records=len(f),primary_models=len(main),source_files=len(sources),spreadsheet_readback=True,main_test_rows=2766),indent=2))
print(main[['model']+MET].to_string(index=False));print('RECORDS',len(f),'MODELS',len(main),'SOURCES',len(sources))
