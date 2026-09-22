"""Export the completed five-method experiment, excluding ESM-LM records."""
import hashlib
import json
import math
from pathlib import Path
import re
import shutil
import zipfile

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
OUT = HERE / "results_no_esm_lm_2697"
RUNS = [HERE / "catapro_test_2697_residual_a3_distinct",
        HERE / "catapro_test_2697_residual_remaining_distinct"]
METHODS = ["A0_pamp", "A1_hotflip", "A2_fw_avg", "A3_fw_avg_pamp", "B2_random"]
ABBR = dict(zip(METHODS, ["SPM", "HF", "FWA", "FWA-P", "RND"]))
ENDPOINTS = ["round1_top1", "round1_best_of_top5", "round2_top1"]
DATASET = Path('/root/kcat-data_0.4simi-10fold.csv')


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2))


def read_csv(path):
    return pd.read_csv(path, float_precision="round_trip")


def mutate(sequence, edits):
    result = list(sequence)
    positions = []
    for edit in edits:
        match = re.fullmatch(r"([A-Z])(\d+)([A-Z])", edit)
        assert match, edit
        old, pos, new = match.groups()
        pos = int(pos) - 1
        assert result[pos] == old and new != old
        result[pos] = new
        positions.append(pos)
    assert len(positions) == len(set(positions))
    assert sum(a != b for a, b in zip(result, sequence)) == len(edits)
    return "".join(result)


def add_sheet(wb, name, frame):
    sheet = wb.create_sheet(name)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(frame.columns))}{len(frame)+1}"
    headers = []
    for i, col in enumerate(frame.columns, 1):
        cell = WriteOnlyCell(sheet, str(col))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="397B69")
        headers.append(cell)
        sheet.column_dimensions[get_column_letter(i)].width = min(44, max(14, len(str(col)) + 2))
    sheet.append(headers)
    for row in frame.itertuples(index=False, name=None):
        cells = []
        for value in row:
            if pd.isna(value):
                value = None
            elif isinstance(value, np.generic):
                value = value.item()
            if isinstance(value, str):
                cell = WriteOnlyCell(sheet, value)
                cell.data_type = "s"
                cells.append(cell)
            else:
                cells.append(value)
        sheet.append(cells)


def main():
    OUT.mkdir(exist_ok=True)
    for run in RUNS:
        assert json.loads((run / "status.json").read_text())["stage"] == "complete"
        assert json.loads((run / "verification.json").read_text())["status"] == "PASS"
    cohorts = [read_csv(run / "cohort.csv") for run in RUNS]
    pd.testing.assert_frame_equal(*cohorts)
    cohort = cohorts[0]
    assert len(cohort) == 2697 and cohort.row_id.is_unique
    ids = set(cohort.row_id)
    contracts = [json.loads((run / "contract.json").read_text()) for run in RUNS]
    for key in ["dataset", "sources", "evaluators", "site_policy", "requested_length_bounds"]:
        assert contracts[0][key] == contracts[1][key]
    for path, digest in contracts[0]["input_hashes"].items():
        if not path.endswith("run.py"):
            assert contracts[1]["input_hashes"][path] == digest
    assert sha(DATASET) == contracts[0]["input_hashes"][str(DATASET)]

    source_metadata = read_csv(DATASET).iloc[cohort.row_id.to_numpy()].copy()
    np.testing.assert_array_equal(source_metadata.Sequence, cohort.sequence_clean)
    source_metadata.insert(0, "row_id", cohort.row_id.to_numpy())
    source_metadata = source_metadata.rename(columns={"Unnamed: 0": "original_dataset_index"})
    source_metadata["sequence_length"] = source_metadata.Sequence.str.len()
    sequences = dict(zip(cohort.row_id, cohort.sequence_clean))
    methods = pd.DataFrame({"method": METHODS, "abbreviation": [ABBR[m] for m in METHODS],
        "display_name": ["Single-point mutation", "HotFlip", "Three-point gradient average",
                         "Three-point gradient average + distance penalty", "Random mutation"],
        "name_zh": ["单点位突变", "HotFlip", "三点梯度平均", "三点梯度平均＋距离惩罚", "随机突变"],
        "implementation": ["One starting-point gradient with distance penalty",
            "One starting-point gradient without distance penalty",
            "Average of three path gradients without distance penalty",
            "Average of three path gradients with distance penalty",
            "Seeded random scores over legal amino-acid substitutions"]})

    candidates = pd.concat([read_csv(run / "candidate_results.csv") for run in RUNS], ignore_index=True)
    candidates = candidates[candidates.method.isin(METHODS)].copy()
    endpoints = pd.concat([read_csv(run / "endpoint_results.csv") for run in RUNS], ignore_index=True)
    endpoints = endpoints[endpoints.method.isin(METHODS)].copy()
    assert len(candidates) == 80910 and len(endpoints) == 40455
    for frame, keys, count in [(candidates, ["row_id", "method", "regime", "candidate_rank"], 6),
                               (endpoints, ["row_id", "method", "endpoint"], 3)]:
        assert frame.source.eq("residual").all() and frame.evaluator.eq("residual").all()
        assert not frame.duplicated(keys).any()
        assert set(frame.method) == set(METHODS)
        assert frame.groupby(["row_id", "method"]).size().eq(count).all()
        assert np.isfinite(frame[["wt_log2", "mutant_log2", "delta_log2"]].to_numpy()).all()
        np.testing.assert_allclose(frame.mutant_log2-frame.wt_log2, frame.delta_log2, atol=1e-12, rtol=1e-12)
        assert frame.groupby("row_id").wt_log2.nunique().eq(1).all()
        frame.insert(frame.columns.get_loc("method") + 1, "method_abbreviation", frame.method.map(ABBR))
        frame["predicted_fold_change"] = np.exp2(frame.delta_log2)
        edit_lists = frame.edits.map(json.loads)
        frame["mutation_1"] = edit_lists.map(lambda x: x[0])
        frame["mutation_2"] = edit_lists.map(lambda x: x[1] if len(x) == 2 else "")
        frame["mutant_sequence"] = [mutate(sequences[int(rid)], edits)
            for rid, edits in zip(frame.row_id, edit_lists)]
        assert np.array_equal(edit_lists.map(len), frame.net_substitutions)
    first = candidates[candidates.regime == "round1_top5"]
    second = candidates[candidates.regime == "round2_top1"]
    assert first.groupby(["row_id", "method"]).candidate_rank.apply(lambda x: set(x) == set(range(1, 6))).all()
    assert first.net_substitutions.eq(1).all() and second.net_substitutions.eq(2).all()
    for endpoint, candidate_frame in [("round1_top1", first[first.candidate_rank == 1]),
            ("round1_best_of_top5", first.sort_values("candidate_rank").loc[
                first.sort_values("candidate_rank").groupby(["row_id", "method"]).mutant_log2.idxmax()]),
            ("round2_top1", second)]:
        left = endpoints[endpoints.endpoint == endpoint].set_index(["row_id", "method"]).sort_index()
        right = candidate_frame.set_index(["row_id", "method"]).sort_index()
        pd.testing.assert_frame_equal(left[candidates.columns.drop(["row_id", "method"])],
                                      right[candidates.columns.drop(["row_id", "method"])])

    summary_rows, trim_rows = [], []
    for method in METHODS:
        for endpoint in ENDPOINTS:
            group = endpoints[(endpoints.method == method) & (endpoints.endpoint == endpoint)]
            assert set(group.row_id) == ids
            x = group.delta_log2.to_numpy()
            q1, med, q3 = np.quantile(x, [.25, .5, .75]); iqr = q3-q1
            inside = x[(x >= q1-1.5*iqr) & (x <= q3+1.5*iqr)]
            summary_rows.append(dict(method=method, method_abbreviation=ABBR[method], endpoint=endpoint,
                N=len(x), mean_delta_log2=x.mean(), geometric_mean_fold=2**x.mean(),
                median_delta_log2=med, median_fold=2**med, std_delta_log2=x.std(ddof=1),
                q1=q1, q3=q3, min=x.min(), max=x.max(), whisker_low=inside.min(), whisker_high=inside.max(),
                positive_fraction=(x>1e-6).mean(), negative_fraction=(x < -1e-6).mean(),
                neutral_fraction=(np.abs(x)<=1e-6).mean()))
            ordered = group.sort_values(["delta_log2", "row_id"])
            for scheme, low, high in [("both_tails_1pct", .01, .01),
                                      ("both_tails_5pct", .05, .05), ("upper_tail_5pct", 0, .05)]:
                lo, hi = math.floor(len(x)*low), math.floor(len(x)*high)
                kept = ordered.iloc[lo:len(x)-hi].delta_log2
                trim_rows.append(dict(method=method, method_abbreviation=ABBR[method], endpoint=endpoint,
                    scheme=scheme, N_before=len(x), removed_lower=lo, removed_upper=hi, N_retained=len(kept),
                    mean_delta_log2=kept.mean(), geometric_mean_fold=2**kept.mean(),
                    median_delta_log2=kept.median(), positive_fraction=(kept>1e-6).mean(),
                    negative_fraction=(kept < -1e-6).mean()))
    summary = pd.DataFrame(summary_rows); trimmed = pd.DataFrame(trim_rows)
    expected = pd.concat([read_csv(run/'summary.csv') for run in RUNS], ignore_index=True)
    expected = expected[expected.method.isin(METHODS)].set_index(["method", "endpoint"]).sort_index()
    actual = summary.set_index(["method", "endpoint"]).sort_index()
    for col in ["N", "mean_delta_log2", "geometric_mean_fold", "median_delta_log2", "positive_fraction", "negative_fraction"]:
        np.testing.assert_allclose(actual[col], expected[col], atol=1e-12, rtol=1e-12)
    matrices = {endpoint: endpoints[endpoints.endpoint == endpoint].pivot(
        index="row_id", columns="method_abbreviation", values="delta_log2").reindex(
            index=cohort.row_id, columns=list(ABBR.values())).reset_index() for endpoint in ENDPOINTS}
    print("Validated all candidates, endpoints, sequences and published summary values.", flush=True)

    tables = {"methods": methods, "cohort": source_metadata, "candidate_results": candidates,
        "endpoint_results": endpoints, "summary": summary, "trimmed_summary": trimmed}
    for name, frame in tables.items():
        frame.to_csv(OUT/f"{name}.csv", index=False, encoding="utf-8-sig", float_format="%.17g")
    for name, frame in matrices.items():
        frame.to_csv(OUT/f"delta_matrix_{name}.csv", index=False, encoding="utf-8-sig", float_format="%.17g")

    # Keep recorded search diagnostics while filtering method-specific choices.
    selection_counts = {"round1": 0, "round2": 0}
    method_round_counts = {(method, phase): 0 for method in METHODS for phase in selection_counts}
    with (OUT/"selection_records.jsonl").open("w") as handle:
        for run in RUNS:
            for path in sorted((run/"selections").glob("*.json")):
                payload = json.loads(path.read_text())
                phase = "round1" if path.stem.endswith("round1") else "round2"
                if phase == "round1":
                    payload["residual"]["methods"] = {k: v for k, v in payload["residual"]["methods"].items() if k in METHODS}
                    choices = payload["residual"]["methods"]
                else:
                    payload["methods"] = {k: v for k, v in payload["methods"].items() if k in METHODS}
                    choices = payload["methods"]
                if not choices:
                    continue
                for method in choices:
                    method_round_counts[(method, phase)] += 1
                selection_counts[phase] += 1
                record = dict(source_run=run.name, source_file=path.name, row_id=int(path.name.split('_')[0]),
                              phase=phase, payload=payload)
                handle.write(json.dumps(record, ensure_ascii=False, separators=(",", ":"))+"\n")
    assert all(n == 2697 for n in method_round_counts.values())

    notes = [
        ("Scope", "2697 matched records; five methods; no ESM-LM result rows; no trimming in main tables."),
        ("Target", "Full frozen Residual Predictor; substrate fixed during mutation search."),
        ("Candidates", "80910 rows: each record and method has five round-1 candidates and one round-2 candidate."),
        ("Endpoints", "40455 rows: round1_top1, round1_best_of_top5, round2_top1 for each record and method."),
        ("Delta", "delta_log2 = mutant_log2 - wt_log2; predicted_fold_change = 2 ** delta_log2."),
        ("Selection", "Round 2 starts from fixed round-1 rank 1, not best-of-five, and changes a distinct site."),
        ("Original IDs", "row_id is the zero-based row position in the original dataset; original_dataset_index is a separate source column."),
        ("Sequences", "cohort.Sequence is the original sequence; mutant_sequence is reconstructed and validated from edits."),
        ("Selection positions", "Positions in methods choices are zero-based; trace.position and blocked_positions_1based use one-based numbering."),
        ("Shared traces", "selection_records.jsonl preserves source-run shared path diagnostics; a shared trace is not a separate trajectory for every method."),
        ("Trim rules", "Per method and endpoint, sort by delta_log2 then row_id; remove floor(N*p) from specified tails; not an error-labeling rule."),
        ("Matched comparison", "Main tables use the same 2697 records; independent trimming can retain different records across methods."),
        ("Labels", "SPM denotes original A0 by user naming; all methods' round-1 candidates are single-residue substitutions."),
        ("Precision", "CSV/JSON preserve double-precision values; Excel has its standard numeric precision limits."),
        ("Interpretation", "Model prediction changes, not measured mutant activity; original legacy substrate feature pipeline retained."),
        ("Figures", "Three panels, English method descriptions, all observations used; displayed y-axis zoom is [-0.55,0.95].")]
    note_table = pd.DataFrame(notes, columns=["item", "description"])
    wb = Workbook(write_only=True)
    add_sheet(wb, "README", note_table)
    for name, frame in tables.items():
        add_sheet(wb, name, frame)
    for endpoint, frame in matrices.items():
        add_sheet(wb, "delta_"+endpoint, frame)
    workbook = OUT/"all_results_no_esm_lm.xlsx"
    wb.save(workbook)
    check = load_workbook(workbook, read_only=True, data_only=True)
    for name, frame in tables.items():
        assert sum(1 for _ in check[name].iter_rows(values_only=True)) == len(frame)+1
    check.close()
    print("Saved and checked Excel workbook.", flush=True)

    figures = OUT/"figures"; figures.mkdir(exist_ok=True)
    for ext in ["png", "pdf", "svg"]:
        src = HERE/"figures_no_esm_lm"/f"attack_boxplots_no_esm_lm_central.{ext}"
        shutil.copy2(src, figures/src.name)
    source = OUT/"source"; source.mkdir(exist_ok=True)
    shutil.copy2(Path(__file__), source/Path(__file__).name)
    shutil.copy2(HERE/"plot_boxplots_no_esm_lm.py", source/"plot_boxplots_no_esm_lm.py")
    source_hashes = {}
    for run in RUNS:
        for name in ["cohort.csv", "candidate_results.csv", "endpoint_results.csv", "contract.json", "summary.csv"]:
            source_hashes[str(run/name)] = sha(run/name)
    source_hashes[str(DATASET)] = sha(DATASET)
    package_contract = {k: v for k, v in contracts[0].items() if k not in ["methods", "input_hashes"]}
    package_contract.update(methods=METHODS, package_records=2697, source_artifact_sha256=source_hashes,
        frozen_experiment_inputs=[c["input_hashes"] for c in contracts], selection_record_counts=selection_counts)
    write_json(OUT/"provenance.json", package_contract)

    report = "# 五种攻击方法完整结果（2697 条记录）\n\n"
    report += "包含 SPM（A0，单点位突变）、HF（A1）、FWA（A2）、FWA-P（A3）、RND（B2）。不包含 ESM-LM 结果行。\n\n"
    report += "## 文件\n\n| 文件 | 内容 |\n|---|---|\n"
    for name, frame in tables.items():
        report += f"| {name}.csv | {len(frame):,} 行 |\n"
    report += "| all_results_no_esm_lm.xlsx | 上述表格、字段说明及三份逐记录增益矩阵 |\n"
    report += "| delta_matrix_*.csv | 三种评价方案下，每条记录 × 五种方法的 Δlog₂ |\n"
    report += "| selection_records.jsonl | 过滤后的原始候选选择和共享路径诊断记录 |\n"
    report += "| figures/ | 最终三面板箱形图，英文图例，PNG/PDF/SVG |\n"
    report += "| provenance.json / checksums.json / verification.json | 数据来源、校验和及完整性检查 |\n\n"
    report += "## 完整队列结果（未截尾）\n\n| 方法 | 首轮 Top1 均值 | Top5 最优均值 | 两轮 Top1 均值 | 两轮中位数 | 两轮提升比例 |\n|---|---:|---:|---:|---:|---:|\n"
    for method in METHODS:
        s = summary[summary.method == method].set_index("endpoint")
        report += f"| {ABBR[method]} | {s.loc['round1_top1','mean_delta_log2']:.6f} | {s.loc['round1_best_of_top5','mean_delta_log2']:.6f} | {s.loc['round2_top1','mean_delta_log2']:.6f} | {s.loc['round2_top1','median_delta_log2']:.6f} | {s.loc['round2_top1','positive_fraction']:.2%} |\n"
    report += "\n均值、中位数的单位为 Δlog₂(kcat)。主要 CSV 和 Excel 明细包含全部观测；trimmed_summary.csv 另给出完整队列的双侧各 1%、双侧各 5%、仅上侧 5% 截尾结果。\n\n"
    report += "## 字段与解释\n\n"
    for name, text in notes:
        report += f"- **{name}**: {text}\n"
    report += "\n完整突变序列已按 edits 重建，并校验首轮一个位点、第二轮两个不同位点。原始酶与底物信息见 cohort.csv，以 row_id 关联。"
    report += "这是一份实验结果交付包，不复制模型权重或大型 embedding 缓存；来源路径与输入哈希记录在 provenance.json。\n"
    (OUT/"README.md").write_text(report)
    # Scan result-bearing tables and selection payloads for excluded method IDs.
    for path in [*OUT.glob("*.csv"), OUT/"selection_records.jsonl"]:
        assert "B0_esm_lm" not in path.read_text(encoding="utf-8-sig")
    verification = dict(status="PASS", records=2697, methods=METHODS, candidate_rows=len(candidates),
        endpoint_rows=len(endpoints), summary_rows=len(summary), trimmed_summary_rows=len(trimmed),
        matched_cohort=True, no_excluded_method_result_rows=True, distinct_sites_verified=True,
        mutant_sequences_verified=True, published_summary_reproduced=True, excel_row_counts_verified=True,
        selection_method_round_counts={f"{m}:{phase}": n for (m, phase), n in method_round_counts.items()})
    write_json(OUT/"verification.json", verification)
    checksums = {str(p.relative_to(OUT)): sha(p) for p in sorted(OUT.rglob("*")) if p.is_file() and p.name != "checksums.json"}
    write_json(OUT/"checksums.json", checksums)
    archive = OUT.with_suffix(".zip")
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as z:
        for p in sorted(OUT.rglob("*")):
            if p.is_file():
                z.write(p, str(p.relative_to(OUT.parent)))
    with zipfile.ZipFile(archive) as z:
        assert z.testzip() is None
    print(json.dumps(dict(status="PASS", directory=str(OUT), zip=str(archive),
        zip_bytes=archive.stat().st_size, excel=str(workbook), files=len(checksums)+1), ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
